import csv
import io
import openpyxl

# maps a bunch of likely column headers to our internal field names
COLUMN_MAP = {
    "english": "english", "verified english": "english", "word": "english",
    "french": "french", "verified french": "french",
    "gpt ewe": "gpt_ewe", "ewe": "gpt_ewe", "ewe translation": "gpt_ewe",
    "gpt twi": "gpt_twi", "twi": "gpt_twi",
    "confidence": "confidence",
    "validated sentiment": "ai_sentiment", "sentiment": "ai_sentiment",
    "validated pos": "ai_pos", "pos": "ai_pos", "part of speech": "ai_pos",
    "comments": "ai_comment", "comment": "ai_comment",
    "id": "source_ref",
}


def _normalise_header(h):
    return (h or "").strip().lower()


def _row_to_entry(header, row):
    entry = {}
    for col_name, value in zip(header, row):
        field = COLUMN_MAP.get(_normalise_header(col_name))
        if field:
            entry[field] = value
    return entry


def parse_lexicon_file(filename, file_bytes):
    """Returns a list of dicts with keys matching Entry model fields.
    Supports .xlsx (first matching sheet, or first sheet) and .csv."""
    name = filename.lower()
    entries = []

    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = None
        for candidate in wb.sheetnames:
            if "verified" in candidate.lower() or "lexicon" in candidate.lower():
                sheet = wb[candidate]
                break
        if sheet is None:
            sheet = wb[wb.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        header_norm = tuple(_normalise_header(h) for h in header)
        for row in rows:
            if row is None or all(v is None for v in row):
                continue
            row_norm = tuple(_normalise_header(v) if isinstance(v, str) else v for v in row)
            if row_norm == header_norm:
                continue  # repeated header row baked into the sheet
            e = _row_to_entry(header, row)
            if e.get("english") or e.get("gpt_ewe"):
                entries.append(e)

    elif name.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        header = next(reader)
        for row in reader:
            if not row or all(not c.strip() for c in row):
                continue
            e = _row_to_entry(header, row)
            if e.get("english") or e.get("gpt_ewe"):
                entries.append(e)
    else:
        raise ValueError("unsupported file type, use .xlsx or .csv")

    return entries
